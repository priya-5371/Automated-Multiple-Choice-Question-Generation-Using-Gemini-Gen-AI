import streamlit as st
import sqlite3
import bcrypt
import google.generativeai as genai
import json
import uuid
from fpdf import FPDF
from gtts import gTTS
from docx import Document
from PyPDF2 import PdfReader
import PyPDF2
import os


# Configure Gemini API key
genai.configure(api_key="AIzaSyAkmsXJIMknAkMHstTNSTI8xAXBTBe5HGs")

# SQLite Database Setup
conn = sqlite3.connect("users.db", check_same_thread=False)
c = conn.cursor()


# Create necessary tables if they don't exist
c.execute('''CREATE TABLE IF NOT EXISTS users (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    username TEXT UNIQUE,
    password TEXT,
    role TEXT
)''')

c.execute('''CREATE TABLE IF NOT EXISTS quizzes (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    teacher TEXT,
    title TEXT,
    mcqs TEXT,
    answers TEXT
)''')

c.execute('''CREATE TABLE IF NOT EXISTS lectures (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    teacher TEXT,
    title TEXT,
    content TEXT
)''')

c.execute('''CREATE TABLE IF NOT EXISTS scores (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    student TEXT,
    quiz_id INTEGER,
    student_answers TEXT,
    score INTEGER DEFAULT NULL,
    evaluated_by TEXT DEFAULT NULL,
    status TEXT DEFAULT 'Pending'
)''')

c.execute('''CREATE TABLE IF NOT EXISTS quiz_titles (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    title TEXT UNIQUE
)''')

c.execute('''CREATE TABLE IF NOT EXISTS marks (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    student_id INTEGER,
    quiz_title TEXT,
    subject TEXT,
    teacher_name TEXT,
    score INTEGER,
    status TEXT DEFAULT 'Pending'
)''')


def ensure_status_column():
    try:
        c.execute("ALTER TABLE scores ADD COLUMN status TEXT DEFAULT 'Pending'")
        conn.commit()
    except sqlite3.OperationalError as e:
        if "duplicate column name" not in str(e).lower():
            st.error(f"Database error: {e}")
    except sqlite3.Error as e:
        st.error(f"Database error: {e}")

ensure_status_column()
conn.commit()

def update_score_and_rerun(student, quiz_id, score):
    try:
        c.execute("UPDATE scores SET score = ?, evaluated_by = ?, status = 'Assigned' WHERE student = ? AND quiz_id = ?",
                  (score, st.session_state.username, student, quiz_id))
        c.execute("UPDATE marks SET score = ?, status='Assigned' WHERE student_id = ? AND quiz_title = ?", (score, student, "Untitled"))
        conn.commit()
        st.success(f"✅ Score submitted for {student} (Quiz {quiz_id})")
        st.experimental_rerun()
    except sqlite3.Error as e:
        st.error(f"Database error: {e}")

def generate_mcqs(content, num_mcqs, difficulty):
    model = genai.GenerativeModel(model_name="gemini-1.5-flash")

    # Define different prompts for difficulty levels
    difficulty_prompts = {
        "Easy": f"Generate {num_mcqs} simple MCQs from the following content. Ensure basic conceptual questions with direct answers:\n\n{content}",
        "Medium": f"Generate {num_mcqs} MCQs from the following content. Include a mix of conceptual and application-based questions with some tricky distractors:\n\n{content}",
        "Hard": f"Generate {num_mcqs} challenging MCQs from the following content. Ensure complex application-based, inference, and analysis questions with strong distractors:\n\n{content}"
    }

    prompt = difficulty_prompts.get(difficulty, difficulty_prompts["Medium"])  # Default to Medium
    response = model.generate_content(prompt)
    
    if response and response.text:
        return response.text.strip()
    return ""



# Function to generate audio for MCQs
def generate_mcq_audio(question, options):
    text = question + " " + " ".join(options)
    tts = gTTS(text=text, lang="en")
    filename = f"mcq_{uuid.uuid4().hex}.mp3"
    tts.save(filename)
    return filename

# Function to extract text from uploaded files
def extract_text_from_file(uploaded_file):
    if uploaded_file.type == "application/pdf":
        pdf_reader = PdfReader(uploaded_file)
        return "\n".join([page.extract_text() for page in pdf_reader.pages if page.extract_text()])
    elif uploaded_file.type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = Document(uploaded_file)
        return "\n".join([para.text for para in doc.paragraphs])
    elif uploaded_file.type == "text/plain":
        return uploaded_file.getvalue().decode("utf-8")
    return ""




# ✅ Function to save MCQs as PDF
def save_mcqs_to_pdf(mcqs_text, title):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)  # ✅ Use built-in font (fixes missing font issue)

    pdf.cell(200, 10, txt=title, ln=True, align='C')
    pdf.multi_cell(0, 10, mcqs_text)

    filename = f"quiz_{uuid.uuid4().hex}.pdf"
    pdf.output(filename, "F").encode("utf-8")  # ✅ Save file
    return filename

# User Registration Function
def register_user(username, password, role):
    hashed_password = bcrypt.hashpw(password.encode(), bcrypt.gensalt())
    try:
        c.execute("INSERT INTO users (username, password, role) VALUES (?, ?, ?)", (username, hashed_password, role))
        conn.commit()
        return True
    except sqlite3.IntegrityError:
        return False

# User Authentication Function
def authenticate_user(username, password):
    c.execute("SELECT password, role FROM users WHERE username = ?", (username,))
    user = c.fetchone()
    if user and bcrypt.checkpw(password.encode(), user[0]):
        return user[1]  
    return None

# Function to parse MCQs safely
def parse_mcqs(mcq_text):
    try:
        mcq_list = mcq_text.strip().split("\n")
        parsed_questions = []
        i = 0

        while i < len(mcq_list):
            question = mcq_list[i].strip()
            if not question or not question[0].isdigit():
                i += 1
                continue

            i += 1
            options = []
            while i < len(mcq_list) and len(options) < 4:
                option = mcq_list[i].strip()
                if option.startswith(("a)", "b)", "c)", "d)")):
                    options.append(option)
                i += 1

            if len(options) == 4:
                parsed_questions.append((question, options))

        return parsed_questions
    except Exception as e:
        st.error(f"⚠ Error parsing MCQs: {e}")
        return []

def evaluate_student_answers(student_answers, correct_answers):
    score = 0
    st.write("Correct Answers:", correct_answers)
    st.write("Student Answers:", student_answers)




    for q, student_answer in student_answers.items():
        if q in correct_answers:
            if student_answer == correct_answers[q]:
                score += 1
                st.write(f"Question {q}:Correct")
            else:
                st.write(f"Question {q}: Incorrect(Expected:{correct_answers[q]},Got:{student_answer})")
    return score


# Streamlit Page Configuration
st.set_page_config(page_title="Smart Quiz For Learners", layout="wide")

st.markdown(
    """
    <style>
    div[data-testid="stTextInput"] {
        width: 300px !important;  /* Adjust width as needed */
        margin: left;
    }
    </style>
    """,
    unsafe_allow_html=True
)
st.title("📝 Smart Quiz For Learners")

# Login & Registration System
if "page" not in st.session_state:
    st.session_state.page = "login"

st.markdown("<h2 style='text-align: center;'> Your Knowledge, Your Score – Let’s Go!</h2>", unsafe_allow_html=True)


if st.session_state.page == "login":
    col1, col2, col3 = st.columns([1, 1, 1])
    with col2:
        username = st.text_input("👤 Username", key="username_input")
        password = st.text_input("🔒 Password", type="password", key="password_input")
        login_button, register_button = st.columns([1, 1])
        with login_button:
            if st.button("Login"):
                if username and password:
                    role = authenticate_user(username, password)
                    if role:
                        st.session_state.username = username
                        st.session_state.role = role
                        st.session_state.page = "dashboard"
                        st.rerun()
                    else:
                        st.error("❌ Invalid username or password.")
                else:
                    st.error("❌ Please enter both username and password.")
        with register_button:
            if st.button("Register"):
                st.session_state.page = "register"
                st.rerun()

elif st.session_state.page == "register":
    col1, col2, col3 = st.columns([1, 1, 1])
    with col2:
        new_user = st.text_input("👤 Username",key="new_user")
        new_pass = st.text_input("🔒 Password",type="password", key="new_pass")
        role = st.selectbox("Role", ["Teacher", "Student"],key="role_select")
        register_button, back_button = st.columns([1, 1])
        with register_button:
            if st.button("Register",key="register_btn"):
                if new_user and new_pass:
                    if register_user(new_user, new_pass, role):
                        st.session_state.registered = True  # ✅ Store success message
                        st.session_state.page = "register"  # Stay on the register page
                        st.rerun()
                    else:
                        st.error("❌ Username already exists.")
                else:
                    st.error("❌ Please enter both username and password.")
        if st.session_state.get("registered"):
            st.success("✅ Registration successful! Please login.")
            
        with back_button:
            if st.button("Back to Login"):
                st.session_state.page = "login"
                st.rerun()

elif st.session_state.page == "dashboard":
    st.sidebar.success(f"✅ Logged in as {st.session_state.username} ({st.session_state.role})")
    if st.sidebar.button("Logout"):
        st.session_state.clear()
        st.session_state.page = "login"
        st.rerun()

    if st.session_state.role == "Teacher":
        st.subheader("📚 Teacher Dashboard")

        # ✅ Upload Lecture
        st.subheader("Upload a Lecture")
        lecture_title = st.text_input("Lecture Title:")
        lecture_content = st.text_area("Lecture Content:")

        if st.button("Upload Lecture"):
            if lecture_title and lecture_content:
                c.execute("INSERT INTO lectures (teacher, title, content) VALUES (?, ?, ?)", 
                          (st.session_state.username, lecture_title, lecture_content))
                conn.commit()
                st.success("✅ Lecture uploaded successfully!")
            else:
                st.error("❌ Please enter both a title and content.")

        # ✅ Generate and Post MCQs
        st.subheader("Create and Post MCQs")
        st.subheader("Upload Lecture  Source (PDF, DOCX, TEXT)")
        quiz_title = st.text_input("Quiz Title:")
        # topic = st.text_input("Enter the Topic:")
        num_mcqs = st.number_input("Number of MCQs to Generate:", min_value=1, max_value=20, value=5)
        uploaded_file = st.file_uploader("Choose a file", type=["pdf", "docx", "txt"])
        manual_input = st.text_area("Or Enter a Topic/Text (if no file uploaded)", "")
        mcq_content = ""
        if uploaded_file:
            extracted_text = extract_text_from_file(uploaded_file)
        elif manual_input.strip():
            extracted_text = manual_input
        else:
            extracted_text = ""


        if extracted_text:
            st.text_area("Extracted Content:", extracted_text, height=200)
            if "mcqs_generated" not in st.session_state:
                st.session_state.mcqs_generated = ""
                st.session_state.quiz_generated = False
                st.session_state.quiz_posted = False
            if "generated_mcqs" not in st.session_state:
                st.session_state.generated_mcqs = ""
            if "current_quiz_title" not in st.session_state:
                st.session_state.current_quiz_title = ""
            if "quiz_posted" not in st.session_state:
                st.session_state.quiz_posted = False
            difficulty = st.selectbox("Select Difficulty Level:", ["Easy", "Medium", "Hard"])
            if st.button("Generate MCQs"):
                if not quiz_title.strip():
                    st.error("Quiz title cannot be empty!")
                elif not extracted_text:  #Check if text was extracted correctly
                    st.error("No text extracted from the file. Please upload a valid file.")
                else:
                    st.session_state.generated_mcqs = generate_mcqs(extracted_text, num_mcqs, difficulty)
                    st.session_state.current_quiz_title = quiz_title
                    st.session_state.selected_difficulty = difficulty
                    st.rerun()
            
            if st.session_state.generated_mcqs:
                st.text_area("✅ MCQs Generated. Please review before posting", st.session_state.generated_mcqs, height=300)
                
                col1, col2 = st.columns(2)
                with col1:
                    if st.button("Regenerate MCQs"):
                        st.session_state.generated_mcqs = generate_mcqs(manual_input, num_mcqs,st.session_state.selected_difficulty)
                        st.rerun()
                with col2:
                    if st.button("Post Quiz"):
                        try:
                            # Check if quiz title already exists
                            c.execute("SELECT 1 FROM quiz_titles WHERE title = ?", (st.session_state.current_quiz_title,))
                            if c.fetchone():
                                st.error("⚠ A quiz with this topic has be posted already.")
                            else:
                                # Insert into quiz_titles first
                                c.execute("INSERT INTO quiz_titles (title) VALUES (?)", (st.session_state.current_quiz_title,))
                                
                                # Insert quiz into quizzes table
                                c.execute("INSERT INTO quizzes (teacher, title, mcqs, answers) VALUES (?, ?, ?, ?)",
                                        (st.session_state.get("username", "Unknown"), 
                                        st.session_state.get("current_quiz_title", "Untitled"), 
                                        st.session_state.get("generated_mcqs", ""), ""))
                                conn.commit()
                                
                                st.session_state.quiz_posted = True  # ✅ Set session state to track post success
                                st.success("✅ Quiz successfully posted!")

                                # Generate and provide a download link for PDF
                                pdf_file = save_mcqs_to_pdf(st.session_state.generated_mcqs, st.session_state.current_quiz_title)
                                with open(pdf_file, "rb") as file:
                                    st.download_button(label="📥 Download Quiz as PDF", data=file, file_name=pdf_file, mime="application/pdf")

                        except sqlite3.IntegrityError:
                            st.error("⚠ A quiz with this topic has be posted already.")
      

            
        # ✅ View & Evaluate Student Submissions
        st.subheader("📊 Evaluate Student Submissions")
        c.execute("SELECT student, quiz_id, student_answers FROM scores WHERE score IS NULL")
        submissions = c.fetchall()

        for student, quiz_id, student_answers in submissions:
            st.markdown(f"### {student} - Quiz {quiz_id}")
            formatted_answers = json.loads(student_answers)

            for q, ans in formatted_answers.items():
                st.write(f"**Q{q}**: {ans}")

            score = st.number_input(f"Assign Score for {student} (Quiz {quiz_id}):", min_value=0, max_value=100, step=1)
            if st.button(f"Submit Score for {student} (Quiz {quiz_id})", key=f"submit_{student}_{quiz_id}"):
                c.execute("UPDATE scores SET score = ?, evaluated_by = ? WHERE student = ? AND quiz_id = ?", 
                          (score, st.session_state.username, student, quiz_id))
                conn.commit()
                st.success(f"✅ Score submitted for {student} (Quiz {quiz_id})")

    elif st.session_state.role == "Student":
        st.subheader("📖 Student Dashboard")


        #View Lectures
        st.subheader("View Lectures")
        c.execute("SELECT title, content FROM lectures")
        lectures = c.fetchall()

        for title, content in lectures:
            st.markdown(f"### {title}")
            st.write(content)
         
        # ✅ Attempt a Quiz
        st.subheader("Attempt a Quiz")
        # c.execute("SELECT id, title, mcqs FROM quizzes")
        c.execute("SELECT id, title, mcqs, answers FROM quizzes")
        quizzes = c.fetchall()

        c.execute("SELECT quiz_title FROM marks WHERE student_id = ?", (st.session_state.username,))
        attempted_quiz_titles = {row[0] for row in c.fetchall()}
        unattempted_quizzes = [q for q in quizzes if q[1] not in attempted_quiz_titles]
        attempted_quizzes = [q for q in quizzes if q[1] in attempted_quiz_titles]

        # Display summary of quizzes
        st.write(f"📊 *Quizzes Summary:*")
        st.write(f"- Total Quizzes: {len(quizzes)}")
        st.write(f"- Attempted Quizzes: {len(attempted_quizzes)}")
        st.write(f"- Unattempted Quizzes: {len(unattempted_quizzes)}")



        if unattempted_quizzes:
            selected_quiz_title = st.selectbox(
                "Select a Quiz to Attempt:",
                [q[1] for q in unattempted_quizzes]
            )
            selected_quiz = next(q for q in unattempted_quizzes if q[1] == selected_quiz_title)
            st.markdown(f"### {selected_quiz[1]}")
            parsed_questions = parse_mcqs(selected_quiz[2])
            student_answers = {}
            for i, (question, options) in enumerate(parsed_questions):
                st.markdown(f"{question}")
                quiz_id_row = c.fetchone()
                quiz_id = quiz_id_row[0] if quiz_id_row else None 
                # audio_file = generate_mcq_audio(question, options)
                # if st.button(f"🔊 Listen to Q{i+1}", key=f"listen_{quiz_id}_{i}"):
                #         st.audio(audio_file, format="audio/mp3")
                key = f"q_{selected_quiz[0]}{i}{st.session_state.username}"
                if key not in st.session_state:
                    st.session_state[key] = options[0]
                selected_option = st.radio(f"Select an option for Q{i+1}:", options, key=key)
                student_answers[i] = selected_option

            c.execute("SELECT id FROM quizzes WHERE id=?", (selected_quiz[0],))
            quiz_id_row = c.fetchone()
            quiz_id = quiz_id_row[0] if quiz_id_row else None  # Extract the actual ID



            if st.button(f"Submit Answers"):
                if len(student_answers) < len(parsed_questions):
                    st.error("⚠ Please answer all questions before submitting!")
                else:
                    if quiz_id is not None:
                            student_response = json.dumps(student_answers)
                            c.execute("REPLACE INTO scores (student, quiz_id, student_answers) VALUES (?, ?, ?)",
                                    (st.session_state.username, quiz_id, student_response))
                            conn.commit()
                            st.success("✅ Answers submitted for evaluation!")
                    else:
                        st.error("⚠ Quiz ID not found!")
                        
        else:
            st.info("🎉 You have attempted all available quizzes!")

 